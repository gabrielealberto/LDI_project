import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from core.ldi_engine import export_ldi_excel


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.result = {
            "portfolio": pd.DataFrame(
                {
                    "isincode": ["TEST"],
                    "description": ["Test bond"],
                    "issuerdescription": ["Test issuer"],
                    "ratingsp": ["AA"],
                    "redemptiondate": ["01/01/2030"],
                    "lots": [1],
                    "nominal_eur": [1000.0],
                    "purchase_value_eur": [990.0],
                    "purchase_commission_eur": [2.95],
                    "sale_commission_eur": [0.0],
                    "cost_eur": [992.95],
                    "maturity_years": [4.0],
                }
            ),
            "cashflow_match": pd.DataFrame(
                {
                    "month": ["2027-01", "2028-01"],
                    "liability_eur": [10.0, 1000.0],
                    "asset_cashflow_eur": [20.0, 1000.0],
                    "external_cash_eur": [0.0, 0.0],
                    "net_cashflow_eur": [10.0, 0.0],
                    "cash_balance_eur": [10.0, 10.0],
                }
            ),
            "uncovered_eur": 0.0,
            "terminal_portfolio_cash_eur": 10.0,
            "annualized_return": 0.02,
            "max_nominal_per_bond": 50_000,
            "max_issuer_weight": 0.4,
            "max_positions": 30,
            "coupon_tax_rate": 0.125,
            "capital_gain_tax_rate": 0.125,
            "tax_breakdown": pd.DataFrame(
                {
                    "year": [2028, 2030],
                    "coupon_taxes_eur": [5.0, 0.0],
                    "capital_gain_taxes_eur": [0.0, 1.25],
                    "total_taxes_eur": [5.0, 1.25],
                }
            ),
            "solver_mip_gap": 0.0,
            "solver_objective": 0.0,
            "status": "Optimal",
            "purchase_commission_eur": 2.95,
            "sale_commission_eur": 0.0,
            "eligible_bonds": 1,
            "total_return_eur": 27.05,
            "roi": 0.0272,
            "weighted_average_maturity_years": 4.0,
        }

    def test_client_report_has_client_sheets_and_charts(self):
        with tempfile.TemporaryDirectory() as directory:
            output = export_ldi_excel(self.result, Path(directory) / "report.xlsx")
            workbook = load_workbook(output, data_only=False)

        self.assertEqual(
            workbook.sheetnames,
            [
                "Client Summary",
                "Portfolio",
                "Cash Flow Profile",
                "Annual Overview",
                "Payed taxes",
                "Methodology & Controls",
            ],
        )
        self.assertEqual(workbook["Client Summary"]["A1"].value, "LDI Portfolio")
        self.assertGreaterEqual(len(workbook["Client Summary"]._charts), 2)
        self.assertGreaterEqual(len(workbook["Annual Overview"]._charts), 2)
        self.assertEqual(workbook["Payed taxes"]["A5"].value, 2028)
        self.assertEqual(workbook["Payed taxes"]["C6"].value, 1.25)
        self.assertNotIn(
            "Selection reason", [cell.value for cell in workbook["Portfolio"][4]]
        )
